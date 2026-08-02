"""Phase 9 — Company Intelligence reports.

Generates:
  reports/generated/company_intelligence_ranked_<date>.xlsx
  reports/generated/company_intelligence_summary_<date>.pdf
  reports/generated/company_resolution_audit_<date>.md
  reports/generated/company_resolution_audit_<date>.xlsx
"""

from __future__ import annotations

import sqlite3
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from pipeline.config.settings import REPORTS_GENERATED_DIR

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_ALT_FILL = PatternFill("solid", fgColor="F2F2F2")


def _today() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d")


def _autosize(ws, headers):
    for i, _ in enumerate(headers, start=1):
        letter = get_column_letter(i)
        longest = max(
            [len(str(headers[i - 1]))]
            + [len(str(c.value)) if c.value is not None else 0 for c in ws[letter]]
        )
        ws.column_dimensions[letter].width = min(max(longest + 2, 10), 55)


def _style_header(ws, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def _ranked_rows(conn: sqlite3.Connection, limit: int = 5000):
    return conn.execute(
        """
        SELECT c.id, c.display_name, c.city, c.state, c.company_type_primary,
               c.license_number,
               ci.company_priority_score, ci.company_priority_tier,
               ci.total_projects, ci.active_projects, ci.projects_last_30_days,
               ci.average_opportunity_score, ci.highest_opportunity_score,
               ci.municipality_count, ci.commercial_project_count,
               ci.residential_project_count, ci.latest_activity_date,
               ci.estimated_opportunity_total, ci.activity_trend
        FROM companies c
        LEFT JOIN company_intelligence ci ON ci.company_id=c.id
        WHERE c.lifecycle_state='active'
        ORDER BY COALESCE(ci.company_priority_score,0) DESC,
                 ci.latest_activity_date DESC
        LIMIT ?
        """,
        (limit,),
    ).fetchall()


def _roles_map(conn: sqlite3.Connection) -> dict[int, str]:
    out: dict[int, list[str]] = {}
    for r in conn.execute("SELECT company_id, role_type, is_primary FROM company_roles "
                          "ORDER BY is_primary DESC"):
        out.setdefault(r["company_id"], []).append(r["role_type"])
    return {cid: ", ".join(roles) for cid, roles in out.items()}


def generate_company_ranked_xlsx(conn: sqlite3.Connection, date: str) -> str:
    rows = _ranked_rows(conn)
    roles = _roles_map(conn)
    wb = Workbook()
    ws = wb.active
    ws.title = "Companies"
    headers = [
        "Rank", "Company", "Roles", "Priority Score", "Priority Tier",
        "Total Projects", "Active Projects", "Projects (30d)",
        "Avg Opportunity Score", "Highest Score", "Latest Activity",
        "Municipality Count", "Commercial", "Residential",
        "Est. Opportunity Total", "Activity Trend", "License", "Data Quality Flags",
    ]
    ws.append(headers)
    for rank, r in enumerate(rows, start=1):
        flags = []
        if not r["license_number"]:
            flags.append("no license")
        if r["company_priority_score"] is None:
            flags.append("no metrics")
        ws.append([
            rank, r["display_name"], roles.get(r["id"], r["company_type_primary"] or ""),
            r["company_priority_score"], r["company_priority_tier"],
            r["total_projects"], r["active_projects"], r["projects_last_30_days"],
            r["average_opportunity_score"], r["highest_opportunity_score"],
            r["latest_activity_date"], r["municipality_count"],
            r["commercial_project_count"], r["residential_project_count"],
            r["estimated_opportunity_total"], r["activity_trend"],
            r["license_number"], ", ".join(flags),
        ])
    for row_idx in range(2, ws.max_row + 1):
        if row_idx % 2 == 0:
            for col in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col).fill = _ALT_FILL
    _style_header(ws, len(headers))
    _autosize(ws, headers)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    path = REPORTS_GENERATED_DIR / f"company_intelligence_ranked_{date}.xlsx"
    REPORTS_GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return str(path)


def generate_company_summary_pdf(conn: sqlite3.Connection, date: str) -> str:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, PageBreak,
    )

    styles = getSampleStyleSheet()
    path = REPORTS_GENERATED_DIR / f"company_intelligence_summary_{date}.pdf"
    REPORTS_GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(path), pagesize=letter,
                            topMargin=0.7 * inch, bottomMargin=0.6 * inch)
    story = []

    total = conn.execute("SELECT COUNT(*) n FROM companies WHERE lifecycle_state='active'").fetchone()["n"]
    pending = conn.execute("SELECT COUNT(*) n FROM company_match_review_queue WHERE lifecycle_state='pending'").fetchone()["n"]
    tier_rows = conn.execute(
        "SELECT company_priority_tier t, COUNT(*) n FROM company_intelligence "
        "GROUP BY company_priority_tier"
    ).fetchall()
    tiers = {r["t"]: r["n"] for r in tier_rows}

    story.append(Paragraph("CorridorIQ — Company Intelligence Summary", styles["Title"]))
    story.append(Paragraph(f"Generated {date}", styles["Normal"]))
    story.append(Spacer(1, 14))
    story.append(Paragraph("Executive summary", styles["Heading2"]))
    story.append(Paragraph(
        f"{total:,} canonical companies derived from real permit records. "
        f"{pending:,} ambiguous identity matches await review. Priority score "
        f"is independent of permit opportunity scoring.", styles["Normal"]))
    story.append(Spacer(1, 10))
    summary_data = [["Priority tier", "Companies"]] + [
        [t, str(tiers.get(t, 0))] for t in ("Critical", "High", "Medium", "Low")
    ]
    st = Table(summary_data, hAlign="LEFT")
    st.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))
    story.append(st)
    story.append(PageBreak())

    story.append(Paragraph("Top 50 companies by priority", styles["Heading2"]))
    rows = _ranked_rows(conn, limit=50)
    roles = _roles_map(conn)
    data = [["#", "Company", "Tier", "Score", "Projects", "Active", "Avg", "Munis"]]
    for i, r in enumerate(rows, start=1):
        data.append([
            str(i), (r["display_name"] or "")[:34],
            r["company_priority_tier"] or "—",
            f"{r['company_priority_score']:.0f}" if r["company_priority_score"] is not None else "—",
            str(r["total_projects"] or 0), str(r["active_projects"] or 0),
            f"{r['average_opportunity_score']:.0f}" if r["average_opportunity_score"] is not None else "—",
            str(r["municipality_count"] or 0),
        ])
    tbl = Table(data, repeatRows=1, hAlign="LEFT",
                colWidths=[0.35*inch, 2.6*inch, 0.7*inch, 0.5*inch, 0.7*inch, 0.6*inch, 0.5*inch, 0.5*inch])
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
    ]))
    story.append(tbl)
    doc.build(story)
    return str(path)


def generate_company_resolution_audit(conn: sqlite3.Connection, date: str) -> tuple[str, str]:
    # Metrics from audit log + queue + link coverage.
    def count(sql, params=()):
        return conn.execute(sql, params).fetchone()["n"]

    processed = count(
        "SELECT COUNT(*) n FROM permits WHERE general_contractor_name IS NOT NULL "
        "AND trim(general_contractor_name)<>''")
    created = count("SELECT COUNT(*) n FROM company_identity_audit_log WHERE action='created'")
    linked_permits = count("SELECT COUNT(*) n FROM permits WHERE contractor_company_id IS NOT NULL")
    linked_projects = count("SELECT COUNT(*) n FROM projects WHERE contractor_company_id IS NOT NULL")
    approved = count("SELECT COUNT(*) n FROM company_identity_audit_log WHERE action='match_approved'")
    rejected = count("SELECT COUNT(*) n FROM company_identity_audit_log WHERE action='match_rejected'")
    review_required = count("SELECT COUNT(*) n FROM company_match_review_queue WHERE lifecycle_state='pending'")
    duplicates = count(
        "SELECT COUNT(*) n FROM (SELECT normalized_name FROM companies "
        "WHERE lifecycle_state='active' GROUP BY normalized_name HAVING COUNT(*)>1)")
    unlinked_permits = count(
        "SELECT COUNT(*) n FROM permits WHERE contractor_company_id IS NULL "
        "AND general_contractor_name IS NOT NULL AND trim(general_contractor_name)<>''")
    unlinked_projects = count(
        "SELECT COUNT(*) n FROM projects WHERE contractor_company_id IS NULL")
    total_companies = count("SELECT COUNT(*) n FROM companies WHERE lifecycle_state='active'")
    exact = linked_permits

    ambiguous = conn.execute(
        "SELECT proposed_company_name, match_confidence, candidate_company_id "
        "FROM company_match_review_queue WHERE lifecycle_state='pending' "
        "ORDER BY match_confidence DESC LIMIT 25"
    ).fetchall()

    # Markdown
    md = [
        f"# Company Resolution Audit — {date}",
        "",
        "Confidence-based, non-destructive linking of permits/projects to canonical",
        "companies. Opportunity scores and raw source values are unchanged.",
        "",
        "## Summary",
        "",
        "| Metric | Count |",
        "|---|---:|",
        f"| Source permits with a company name | {processed:,} |",
        f"| Canonical companies (active) | {total_companies:,} |",
        f"| New companies created | {created:,} |",
        f"| Permits linked to a company | {linked_permits:,} |",
        f"| Projects linked to a company | {linked_projects:,} |",
        f"| Review approvals | {approved:,} |",
        f"| Review rejections | {rejected:,} |",
        f"| Review-required (pending) | {review_required:,} |",
        f"| Duplicate-name candidate groups | {duplicates:,} |",
        f"| Unlinked permits (have a name) | {unlinked_permits:,} |",
        f"| Unlinked projects | {unlinked_projects:,} |",
        "",
        "## Top ambiguous names (pending review)",
        "",
        "| Proposed name | Confidence | Candidate company id |",
        "|---|---:|---:|",
    ]
    for a in ambiguous:
        md.append(f"| {a['proposed_company_name']} | {a['match_confidence']:.0f} | "
                  f"{a['candidate_company_id'] or '—'} |")
    if not ambiguous:
        md.append("| _none_ | | |")
    md_path = REPORTS_GENERATED_DIR / f"company_resolution_audit_{date}.md"
    REPORTS_GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    md_path.write_text("\n".join(md), encoding="utf-8")

    # Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric", "Count"])
    for label, value in [
        ("Source permits with a company name", processed),
        ("Canonical companies (active)", total_companies),
        ("New companies created", created),
        ("Permits linked to a company", linked_permits),
        ("Projects linked to a company", linked_projects),
        ("Review approvals", approved),
        ("Review rejections", rejected),
        ("Review-required (pending)", review_required),
        ("Duplicate-name candidate groups", duplicates),
        ("Unlinked permits (have a name)", unlinked_permits),
        ("Unlinked projects", unlinked_projects),
    ]:
        ws.append([label, value])
    _style_header(ws, 2)
    _autosize(ws, ["Metric", "Count"])

    ws2 = wb.create_sheet("Ambiguous")
    amb_headers = ["Proposed name", "Confidence", "Candidate company id"]
    ws2.append(amb_headers)
    for a in ambiguous:
        ws2.append([a["proposed_company_name"], a["match_confidence"], a["candidate_company_id"]])
    _style_header(ws2, len(amb_headers))
    _autosize(ws2, amb_headers)

    xlsx_path = REPORTS_GENERATED_DIR / f"company_resolution_audit_{date}.xlsx"
    wb.save(xlsx_path)
    return str(md_path), str(xlsx_path)


def generate_company_reports(conn: sqlite3.Connection) -> dict:
    date = _today()
    ranked = generate_company_ranked_xlsx(conn, date)
    summary = generate_company_summary_pdf(conn, date)
    audit_md, audit_xlsx = generate_company_resolution_audit(conn, date)
    return {
        "ranked_xlsx": ranked,
        "summary_pdf": summary,
        "audit_md": audit_md,
        "audit_xlsx": audit_xlsx,
    }


if __name__ == "__main__":
    from pipeline.db.database import init_db

    connection = init_db()
    out = generate_company_reports(connection)
    for k, v in out.items():
        print(f"{k}: {v}")
    connection.close()
