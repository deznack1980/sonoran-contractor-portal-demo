"""Formatted Excel export for ranked permit intelligence reports."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from pipeline.config.settings import REPORTS_GENERATED_DIR
from pipeline.reports.report_presentation import (
    lifecycle_fields_for_row,
    priority_badge,
    product_categories_for_row,
    reason_for_score,
    recommended_next_action,
)
from pipeline.reports.report_templates import contractor_display


HEADERS = [
    "Rank",
    "Overall Score (%)",
    "Priority Badge",
    "Lifecycle Stage",
    "Opportunity Timing",
    "Opportunity Date",
    "Days Since Opportunity",
    "Current Status",
    "Permit Number",
    "Permit Type",
    "Issue Date",
    "Contractor Name",
    "Property Owner",
    "Job Address",
    "City",
    "Description of Work",
    "Estimated Material Opportunity",
    "Likely Product Categories",
    "Reason for Score",
    "Recommended Next Action",
    "Recommended Sales Action",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ALT_FILL = PatternFill("solid", fgColor="F2F2F2")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
YELLOW_FILL = PatternFill("solid", fgColor="FFEB9C")
ORANGE_FILL = PatternFill("solid", fgColor="FCE4D6")
GRAY_FILL = PatternFill("solid", fgColor="D9D9D9")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def _score_fill(score: float | None) -> PatternFill | None:
    if score is None:
        return None
    if score >= 90:
        return GREEN_FILL
    if score >= 80:
        return YELLOW_FILL
    if score >= 70:
        return ORANGE_FILL
    if score >= 60:
        return GRAY_FILL
    return None


def _row_get(row, key: str, default=None):
    try:
        value = row[key]
    except (KeyError, IndexError, TypeError):
        return default
    return default if value is None else value


def _autosize_columns(ws) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[letter]:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 60))
        ws.column_dimensions[letter].width = max(12, max_len + 2)


def build_ranked_rows(db_rows: list) -> list[list]:
    """Convert SQLite project/permit rows into Excel-ready lists."""
    excel_rows: list[list] = []
    for rank, row in enumerate(db_rows, start=1):
        score = _row_get(row, "opportunity_score")
        contractor = contractor_display(row)
        has_contractor = contractor not in ("Not in source",) and not str(contractor).startswith(
            "Project ·"
        )
        description = _row_get(row, "description") or _row_get(row, "project_description") or ""
        permit_row = {
            "permit_type": _row_get(row, "permit_type"),
            "status": _row_get(row, "status"),
            "issued_date": _row_get(row, "issued_date"),
            "filed_date": _row_get(row, "filed_date"),
            "valuation": _row_get(row, "valuation"),
            "square_footage": _row_get(row, "square_footage"),
            "description": description,
        }
        category = _row_get(row, "project_category") or "Other"
        products = ", ".join(product_categories_for_row(row))
        material = _row_get(row, "estimated_material_value")
        life = lifecycle_fields_for_row(row)
        excel_rows.append(
            [
                rank,
                round(float(score), 1) if score is not None else None,
                priority_badge(score),
                life["lifecycle_stage"],
                life["opportunity_timing"],
                life["opportunity_date"] or "",
                life["days_since_opportunity"] if life["days_since_opportunity"] is not None else "",
                _row_get(row, "status") or "",
                _row_get(row, "permit_number") or "",
                _row_get(row, "permit_type") or "",
                _row_get(row, "issued_date") or "",
                contractor,
                (_row_get(row, "owner_name") or "").strip(),
                _row_get(row, "job_address") or "",
                _row_get(row, "city") or "",
                description,
                round(float(material), 2) if material is not None else None,
                products,
                reason_for_score(permit_row, category),
                recommended_next_action(score, has_contractor),
                life["sales_action"],
            ]
        )
    return excel_rows


def write_ranked_permits_excel(
    db_rows: list,
    *,
    filename: str | None = None,
    sheet_title: str = "Ranked Permits",
) -> str:
    """Write a formatted workbook for the ranked permit list."""
    REPORTS_GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    if not filename:
        date_str = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        filename = f"permit_intelligence_ranked_{date_str}.xlsx"
    path = Path(REPORTS_GENERATED_DIR) / filename

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = THIN

    excel_rows = build_ranked_rows(db_rows)
    for r_idx, values in enumerate(excel_rows, start=2):
        for c_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.border = THIN
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if r_idx % 2 == 0:
                cell.fill = ALT_FILL
            if c_idx == 2:
                fill = _score_fill(value if isinstance(value, (int, float)) else None)
                if fill is not None:
                    cell.fill = fill
                    cell.font = Font(bold=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(1, len(excel_rows) + 1)}"
    _autosize_columns(ws)
    ws.row_dimensions[1].height = 22
    wb.save(path)
    return str(path)
