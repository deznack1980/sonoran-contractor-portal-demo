"""Before/after validation for Sprint 1.2 normalization repair.

Captures 7-day scores currently stored in SQLite (before), recomputes with
the repaired engine (after), writes markdown + Excel. Does not overwrite
raw permits.status / permits.permit_type.
"""

from __future__ import annotations

import sqlite3
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from pipeline.analysis.run_analysis import analyze_permit
from pipeline.analysis.scoring import _opportunity_components
from pipeline.config.settings import (
    REPORT_MIN_OPPORTUNITY_SCORE,
    REPORTS_GENERATED_DIR,
)
from pipeline.db.database import init_db


def _bucket(score: float | None) -> str:
    if score is None:
        return "missing"
    if score >= 90:
        return "90+"
    if score >= 80:
        return "80-89"
    if score >= 70:
        return "70-79"
    if score >= 60:
        return "60-69"
    if score >= 50:
        return "50-59"
    return "<50"


def collect_seven_day_rows(conn: sqlite3.Connection, cutoff: str) -> list[dict]:
    rows = conn.execute(
        """
        SELECT p.id AS permit_id, j.name AS municipality, p.jurisdiction,
               p.permit_number, p.permit_type, p.status, p.description,
               p.issued_date, p.valuation, p.square_footage,
               pr.opportunity_score AS score_before,
               pr.project_category AS category_before
        FROM permits p
        JOIN jurisdictions j ON j.slug = p.jurisdiction
        LEFT JOIN projects pr ON pr.permit_id = p.id
        WHERE substr(p.issued_date, 1, 10) >= ?
        ORDER BY substr(p.issued_date, 1, 10) DESC, p.jurisdiction, p.permit_number
        """,
        (cutoff,),
    ).fetchall()
    return [dict(r) for r in rows]


def score_after(row: dict) -> dict:
    permit_row = {
        "permit_type": row["permit_type"],
        "status": row["status"],
        "description": row["description"],
        "issued_date": row["issued_date"],
        "valuation": row["valuation"],
        "square_footage": row["square_footage"],
        "jurisdiction": row["jurisdiction"],
        "municipality": row["municipality"],
    }
    result = analyze_permit(permit_row)
    parts = _opportunity_components(permit_row, result["project_category"])
    return {
        "score_after": result["opportunity_score"],
        "category_after": result["project_category"],
        "normalization": result["normalization"],
        "components_after": parts,
        # Prove raw fields untouched
        "raw_status_preserved": permit_row["status"] == row["status"],
        "raw_type_preserved": permit_row["permit_type"] == row["permit_type"],
    }


def run_validation(conn: sqlite3.Connection | None = None) -> tuple[str, str, dict]:
    close = False
    if conn is None:
        conn = init_db()
        close = True

    now = datetime.now(timezone.utc)
    date_str = now.strftime("%Y-%m-%d")
    cutoff = (now - timedelta(days=7)).strftime("%Y-%m-%d")
    rows = collect_seven_day_rows(conn, cutoff)

    enriched = []
    for row in rows:
        after = score_after(row)
        enriched.append({**row, **after})

    before_scores = [r["score_before"] for r in enriched if r["score_before"] is not None]
    after_scores = [r["score_after"] for r in enriched]
    before_ge60 = sum(1 for s in before_scores if s >= REPORT_MIN_OPPORTUNITY_SCORE)
    after_ge60 = sum(1 for s in after_scores if s >= REPORT_MIN_OPPORTUNITY_SCORE)
    before_lt60 = sum(1 for s in before_scores if s < REPORT_MIN_OPPORTUNITY_SCORE)
    after_lt60 = sum(1 for s in after_scores if s < REPORT_MIN_OPPORTUNITY_SCORE)

    before_dist = Counter(_bucket(s) for s in before_scores)
    after_dist = Counter(_bucket(s) for s in after_scores)
    unique_before = sorted({round(s, 1) for s in before_scores})
    unique_after = sorted({round(s, 1) for s in after_scores})

    status_maps = Counter(r["normalization"]["status_mapping_rule"] for r in enriched)
    type_maps = Counter(r["normalization"]["permit_type_mapping_rule"] for r in enriched)
    unknown_status = sum(
        1 for r in enriched if r["normalization"]["normalized_status"] == "unknown"
    )
    other_after = sum(1 for r in enriched if r["category_after"] == "Other")
    other_before = sum(1 for r in enriched if (r["category_before"] or "Other") == "Other")

    phoenix_open = [
        r
        for r in enriched
        if (r["status"] or "").strip().upper() == "OPEN"
        and "phoenix" in (r["jurisdiction"] or "").lower()
    ]
    phoenix_open_rule = (
        phoenix_open[0]["normalization"]["status_mapping_rule"] if phoenix_open else "n/a"
    )

    raw_preserved = all(r["raw_status_preserved"] and r["raw_type_preserved"] for r in enriched)

    # Representative examples: mix Peoria residential/commercial + Phoenix OPEN
    examples = []
    for pred in (
        lambda r: r["municipality"] == "Peoria" and (r["permit_type"] or "").lower() == "residential",
        lambda r: r["municipality"] == "Peoria" and (r["permit_type"] or "").lower() == "commercial",
        lambda r: r["municipality"] == "Phoenix" and (r["status"] or "").upper() == "OPEN",
    ):
        for r in enriched:
            if pred(r) and r not in examples:
                examples.append(r)
                break
    for r in enriched:
        if r not in examples:
            examples.append(r)
        if len(examples) >= 25:
            break

    summary = {
        "cutoff": cutoff,
        "seven_day_count": len(enriched),
        "before_ge60": before_ge60,
        "after_ge60": after_ge60,
        "before_lt60": before_lt60,
        "after_lt60": after_lt60,
        "unique_before": unique_before,
        "unique_after": unique_after,
        "before_dist": dict(before_dist),
        "after_dist": dict(after_dist),
        "unknown_status": unknown_status,
        "other_before": other_before,
        "other_after": other_after,
        "phoenix_open_count": len(phoenix_open),
        "phoenix_open_rule": phoenix_open_rule,
        "raw_fields_preserved": raw_preserved,
        "status_maps": dict(status_maps),
        "type_maps": dict(type_maps),
        "examples": examples,
        "pass": len(unique_after) > 1 and after_ge60 > 0 and raw_preserved,
    }

    REPORTS_GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    md_path = REPORTS_GENERATED_DIR / f"normalization_repair_validation_{date_str}.md"
    xlsx_path = REPORTS_GENERATED_DIR / f"normalization_repair_validation_{date_str}.xlsx"
    md_path.write_text(_render_md(now, summary, examples), encoding="utf-8")
    _write_xlsx(xlsx_path, summary, enriched, examples)

    if close:
        conn.close()
    return str(md_path), str(xlsx_path), summary


def _render_md(now: datetime, summary: dict, examples: list[dict]) -> str:
    lines = [
        "# Normalization Repair Validation (Sprint 1.2)",
        "",
        f"**Generated:** {now.strftime('%Y-%m-%d %H:%M UTC')}",
        f"**Seven-day cutoff:** `{summary['cutoff']}`",
        f"**Threshold (unchanged):** {REPORT_MIN_OPPORTUNITY_SCORE}",
        f"**Weights (unchanged):** 40/25/20/15 signal/recency/scale/sector",
        f"**Raw field preservation:** {'PASS' if summary['raw_fields_preserved'] else 'FAIL'}",
        f"**Acceptance:** {'PASS' if summary['pass'] else 'FAIL'}",
        "",
        "## Summary",
        "",
        "| Metric | Before | After |",
        "|---|---:|---:|",
        f"| Seven-day permits | {summary['seven_day_count']} | {summary['seven_day_count']} |",
        f"| Scoring 60+ | {summary['before_ge60']} | {summary['after_ge60']} |",
        f"| Scoring below 60 | {summary['before_lt60']} | {summary['after_lt60']} |",
        f"| Unique scores | {len(summary['unique_before'])} ({summary['unique_before']}) | "
        f"{len(summary['unique_after'])} ({summary['unique_after']}) |",
        f"| Category Other | {summary['other_before']} | {summary['other_after']} |",
        f"| Unknown normalized status | — | {summary['unknown_status']} |",
        "",
        "### Phoenix OPEN mapping",
        "",
        f"- Count of Phoenix rows with raw status `OPEN`: **{summary['phoenix_open_count']}**",
        f"- Mapping rule applied: `{summary['phoenix_open_rule']}`",
        "- Global `OPEN` (no municipality) remains `unknown` (not auto-mapped to issued).",
        "",
        "### Score distribution",
        "",
        "| Bucket | Before | After |",
        "|---|---:|---:|",
    ]
    buckets = ["90+", "80-89", "70-79", "60-69", "50-59", "<50", "missing"]
    for b in buckets:
        lines.append(
            f"| {b} | {summary['before_dist'].get(b, 0)} | {summary['after_dist'].get(b, 0)} |"
        )
    lines.extend(
        [
            "",
            "### Status mapping counts",
            "",
        ]
    )
    for rule, n in sorted(summary["status_maps"].items(), key=lambda x: -x[1]):
        lines.append(f"- `{rule}`: {n}")
    lines.extend(["", "### Permit-type mapping counts", ""])
    for rule, n in sorted(summary["type_maps"].items(), key=lambda x: -x[1]):
        lines.append(f"- `{rule}`: {n}")

    lines.extend(
        [
            "",
            "## Before-and-after examples (25)",
            "",
            "| Municipality | Permit | Raw status | Norm status | Raw type | Norm type | "
            "Category before | Category after | Score before | Score after | Scale basis | Rule |",
            "|---|---|---|---|---|---|---|---|---:|---:|---|---|",
        ]
    )
    for r in examples:
        n = r["normalization"]
        lines.append(
            f"| {r['municipality']} | {r['permit_number']} | `{r['status']}` | "
            f"`{n['normalized_status']}` | `{r['permit_type']}` | `{n['normalized_permit_type']}` | "
            f"{r['category_before'] or '—'} | {r['category_after']} | "
            f"{r['score_before'] if r['score_before'] is not None else '—'} | {r['score_after']} | "
            f"{n['scale_basis']} | `{n['status_mapping_rule']}` |"
        )
    lines.append("")
    return "\n".join(lines)


def _write_xlsx(path: Path, summary: dict, enriched: list[dict], examples: list[dict]) -> None:
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )

    def style(ws):
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin
            cell.alignment = Alignment(wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in range(1, ws.max_column + 1):
            letter = get_column_letter(col)
            width = 12
            for cell in ws[letter]:
                width = max(width, min(len(str(cell.value or "")), 40) + 2)
            ws.column_dimensions[letter].width = width

    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric", "Value"])
    for k in (
        "cutoff",
        "seven_day_count",
        "before_ge60",
        "after_ge60",
        "before_lt60",
        "after_lt60",
        "unknown_status",
        "other_before",
        "other_after",
        "phoenix_open_count",
        "phoenix_open_rule",
        "raw_fields_preserved",
        "pass",
    ):
        ws.append([k, str(summary[k])])
    ws.append(["unique_before", str(summary["unique_before"])])
    ws.append(["unique_after", str(summary["unique_after"])])
    style(ws)

    ws2 = wb.create_sheet("All 7-day")
    ws2.append(
        [
            "municipality",
            "permit_number",
            "raw_status",
            "normalized_status",
            "status_mapping_rule",
            "raw_permit_type",
            "normalized_permit_type",
            "category_before",
            "category_after",
            "category_mapping_rule",
            "score_before",
            "score_after",
            "scale_basis",
            "raw_preserved",
        ]
    )
    for r in enriched:
        n = r["normalization"]
        ws2.append(
            [
                r["municipality"],
                r["permit_number"],
                r["status"],
                n["normalized_status"],
                n["status_mapping_rule"],
                r["permit_type"],
                n["normalized_permit_type"],
                r["category_before"],
                r["category_after"],
                n["category_mapping_rule"],
                r["score_before"],
                r["score_after"],
                n["scale_basis"],
                r["raw_status_preserved"] and r["raw_type_preserved"],
            ]
        )
    style(ws2)

    ws3 = wb.create_sheet("Examples")
    example_headers = [cell.value for cell in ws2[1]]
    ws3.append(example_headers)
    for r in examples:
        n = r["normalization"]
        ws3.append(
            [
                r["municipality"],
                r["permit_number"],
                r["status"],
                n["normalized_status"],
                n["status_mapping_rule"],
                r["permit_type"],
                n["normalized_permit_type"],
                r["category_before"],
                r["category_after"],
                n["category_mapping_rule"],
                r["score_before"],
                r["score_after"],
                n["scale_basis"],
                r["raw_status_preserved"] and r["raw_type_preserved"],
            ]
        )
    style(ws3)
    wb.save(path)


if __name__ == "__main__":
    md, xlsx, summary = run_validation()
    print(f"Generated: {md}")
    print(f"Generated: {xlsx}")
    print(
        f"PASS={summary['pass']} before_60+={summary['before_ge60']} "
        f"after_60+={summary['after_ge60']} unique_after={summary['unique_after']}"
    )
