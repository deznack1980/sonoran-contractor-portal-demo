"""Sprint 3 Phase 7 — submitted-vs-issued lead time validation.

Measures how much earlier CorridorIQ can identify projects by tracking
submitted (application) dates instead of waiting for issued permits.

Read-only. Generates:
    reports/generated/lifecycle_lead_time_validation_<date>.md
    reports/generated/lifecycle_lead_time_validation_<date>.xlsx
"""

from __future__ import annotations

import sqlite3
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from pipeline.config.settings import REPORTS_GENERATED_DIR


def _parse(value) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _jurisdiction_stats(conn: sqlite3.Connection, slug: str) -> dict:
    rows = conn.execute(
        "SELECT filed_date, issued_date FROM permits WHERE jurisdiction=?", (slug,)
    ).fetchall()
    total = len(rows)
    submitted = issued = both = missing_submitted = missing_issued = 0
    lead_days: list[int] = []
    for r in rows:
        f = _parse(r["filed_date"])
        i = _parse(r["issued_date"])
        if f:
            submitted += 1
        else:
            missing_submitted += 1
        if i:
            issued += 1
        else:
            missing_issued += 1
        if f and i:
            both += 1
            delta = (i - f).days
            if delta >= 0:
                lead_days.append(delta)
    avg_lead = round(sum(lead_days) / len(lead_days), 1) if lead_days else None
    return {
        "slug": slug,
        "total": total,
        "submitted": submitted,
        "issued": issued,
        "both": both,
        "missing_submitted": missing_submitted,
        "missing_issued": missing_issued,
        "avg_days_submitted_to_issued": avg_lead,
        "submitted_only": submitted - both,  # known via submitted, not yet issued
    }


def generate_lifecycle_lead_time_validation(conn: sqlite3.Connection) -> list[str]:
    jurisdictions = [
        dict(r)
        for r in conn.execute(
            "SELECT slug, name, status FROM jurisdictions "
            "WHERE status='connected' ORDER BY name"
        )
    ]
    stats = []
    for j in jurisdictions:
        s = _jurisdiction_stats(conn, j["slug"])
        s["name"] = j["name"]
        stats.append(s)

    # Overall lead time (permits with both dates), weighted across jurisdictions.
    total_both = sum(s["both"] for s in stats)
    weighted_lead = (
        round(
            sum(
                (s["avg_days_submitted_to_issued"] or 0) * s["both"]
                for s in stats
            )
            / total_both,
            1,
        )
        if total_both
        else None
    )
    total_submitted_only = sum(s["submitted_only"] for s in stats)

    today = date.today().isoformat()
    REPORTS_GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    md_path = REPORTS_GENERATED_DIR / f"lifecycle_lead_time_validation_{today}.md"
    xlsx_path = REPORTS_GENERATED_DIR / f"lifecycle_lead_time_validation_{today}.xlsx"

    # ---------- Markdown ----------
    lines = [f"# Lifecycle Lead-Time Validation — {today}", ""]
    if weighted_lead is not None:
        lines.append(
            f"**Bottom line:** Tracking submitted (application) permits lets "
            f"CorridorIQ identify a project on average **{weighted_lead} days "
            f"earlier** than waiting for the issued permit. "
            f"**{total_submitted_only:,}** projects are currently known via a "
            f"submitted date but have not been issued yet — pure early-lead "
            f"opportunities that issued-only tracking would miss."
        )
    else:
        lines.append(
            "**Bottom line:** No jurisdiction currently supplies both submitted "
            "and issued dates, so lead-time gain cannot be measured. CorridorIQ "
            "falls back to issued dates for `opportunity_date`."
        )
    lines.append("")
    lines.append(
        "| Jurisdiction | Submitted | Issued | Both | Missing submitted | "
        "Missing issued | Avg days submitted→issued | Submitted-only (early leads) |"
    )
    lines.append("| --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: |")
    for s in stats:
        lines.append(
            f"| {s['name']} | {s['submitted']:,} | {s['issued']:,} | {s['both']:,} "
            f"| {s['missing_submitted']:,} | {s['missing_issued']:,} "
            f"| {s['avg_days_submitted_to_issued'] if s['avg_days_submitted_to_issued'] is not None else '—'} "
            f"| {s['submitted_only']:,} |"
        )
    lines.append(
        f"| **Total** | "
        f"{sum(s['submitted'] for s in stats):,} | "
        f"{sum(s['issued'] for s in stats):,} | {total_both:,} | "
        f"{sum(s['missing_submitted'] for s in stats):,} | "
        f"{sum(s['missing_issued'] for s in stats):,} | "
        f"{weighted_lead if weighted_lead is not None else '—'} | "
        f"{total_submitted_only:,} |"
    )
    lines.append("")
    lines.append(
        "> *Avg days submitted→issued* is the concrete lead time gained: how many "
        "days before the issued permit CorridorIQ already had the project via its "
        "application date. Jurisdictions with no submitted date default to issued."
    )
    lines.append("")
    md_path.write_text("\n".join(lines), encoding="utf-8")

    # ---------- Excel ----------
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")

    ws = wb.active
    ws.title = "Lead time by jurisdiction"
    ws.append(
        ["Jurisdiction", "Total", "Submitted", "Issued", "Both",
         "Missing submitted", "Missing issued",
         "Avg days submitted→issued", "Submitted-only (early leads)"]
    )
    for s in stats:
        ws.append(
            [s["name"], s["total"], s["submitted"], s["issued"], s["both"],
             s["missing_submitted"], s["missing_issued"],
             s["avg_days_submitted_to_issued"], s["submitted_only"]]
        )
    ws.append(
        ["TOTAL",
         sum(s["total"] for s in stats),
         sum(s["submitted"] for s in stats),
         sum(s["issued"] for s in stats),
         total_both,
         sum(s["missing_submitted"] for s in stats),
         sum(s["missing_issued"] for s in stats),
         weighted_lead,
         total_submitted_only]
    )
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws.freeze_panes = "A2"
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 12), 40)

    wb.save(xlsx_path)
    return [str(md_path), str(xlsx_path)]


if __name__ == "__main__":
    from pipeline.db.database import init_db

    connection = init_db()
    for p in generate_lifecycle_lead_time_validation(connection):
        print(f"Wrote {p}")
    connection.close()
