"""Sprint 6 reports — supplier catalog import + price comparison.

    python -m pipeline.products.reports

Generates in reports/generated/:
  supplier_catalog_import_<date>.md
  supplier_catalog_import_<date>.xlsx
  supplier_price_comparison_<date>.xlsx

No supplier cost fields are written to the price-comparison report.
"""

from __future__ import annotations

import sqlite3
from collections import Counter
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from pipeline.config import settings
from pipeline.db.database import get_connection
from pipeline.products.routing import compare_offers

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _today() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d")


def _style_header(ws):
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT


def _autosize(ws):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 60)


# --------------------------------------------------------------------------
# Data gathering
# --------------------------------------------------------------------------

def _coverage(conn: sqlite3.Connection) -> dict:
    total_products = conn.execute("SELECT COUNT(*) n FROM products").fetchone()["n"]
    total_offers = conn.execute("SELECT COUNT(*) n FROM supplier_products").fetchone()["n"]
    missing_mpn = conn.execute(
        "SELECT COUNT(*) n FROM products WHERE manufacturer_part_number IS NULL "
        "OR manufacturer_part_number=''").fetchone()["n"]
    per_supplier = conn.execute(
        "SELECT s.name, s.code, COUNT(*) n, "
        "SUM(CASE WHEN sp.quantity_available>0 THEN 1 ELSE 0 END) in_stock "
        "FROM supplier_products sp JOIN suppliers s ON s.id=sp.supplier_id "
        "GROUP BY sp.supplier_id ORDER BY n DESC").fetchall()
    return {
        "total_products": total_products,
        "total_offers": total_offers,
        "missing_mpn": missing_mpn,
        "per_supplier": [dict(r) for r in per_supplier],
    }


def _latest_imports(conn: sqlite3.Connection, limit: int = 20) -> list[dict]:
    rows = conn.execute(
        "SELECT il.*, s.name AS supplier_name FROM supplier_import_log il "
        "LEFT JOIN suppliers s ON s.id=il.supplier_id ORDER BY il.id DESC LIMIT ?",
        (limit,)).fetchall()
    return [dict(r) for r in rows]


def _price_comparison(conn: sqlite3.Connection, jobsite: dict | None = None,
                      quantity: float = 1) -> dict:
    jobsite = jobsite or {"city": "Phoenix", "state": "AZ"}
    product_ids = [r["id"] for r in conn.execute(
        "SELECT DISTINCT p.id FROM products p "
        "JOIN supplier_products sp ON sp.product_id=p.id")]
    cheapest_freq = Counter()
    missing_price_products = 0
    no_available_products = 0
    multi_supplier_products = 0
    rows = []
    for pid in product_ids:
        result = compare_offers(conn, pid, quantity, jobsite, include_cost=False)
        offers = result.get("offers", [])
        if len([o for o in offers]) > 1:
            multi_supplier_products += 1
        if any("missing_price" in o["warnings"] for o in offers):
            missing_price_products += 1
        rec = result.get("recommended_supplier_id")
        if rec is None:
            no_available_products += 1
        else:
            name = next((o["supplier"] for o in offers if o["supplier_id"] == rec), str(rec))
            cheapest_freq[name] += 1
            top = next(o for o in offers if o["supplier_id"] == rec)
            rows.append({
                "product": result["product"]["product_name"],
                "sku": result["product"]["sku"],
                "cheapest_supplier": name,
                "material_subtotal": top["material_subtotal"],
                "delivery_cost": top["estimated_delivery_cost"],
                "estimated_total": top["estimated_total"],
                "suppliers_compared": len(offers),
            })
    return {
        "jobsite": jobsite, "quantity": quantity,
        "products_evaluated": len(product_ids),
        "multi_supplier_products": multi_supplier_products,
        "missing_price_products": missing_price_products,
        "no_available_products": no_available_products,
        "cheapest_frequency": dict(cheapest_freq.most_common()),
        "rows": rows,
    }


# --------------------------------------------------------------------------
# Writers
# --------------------------------------------------------------------------

def _write_import_md(path, coverage, imports):
    L = ["# CorridorIQ — Supplier Catalog Import",
         f"\n_Generated {_today()}_\n",
         "## Catalog coverage\n",
         "| Metric | Value |", "|---|---:|",
         f"| Master products | {coverage['total_products']} |",
         f"| Supplier offers | {coverage['total_offers']} |",
         f"| Products missing mfr part # | {coverage['missing_mpn']} |",
         "\n## Supplier coverage\n",
         "| Supplier | Code | Offers | In stock |", "|---|---|---:|---:|"]
    for s in coverage["per_supplier"]:
        L.append(f"| {s['name']} | {s['code'] or ''} | {s['n']} | {s['in_stock'] or 0} |")

    L.append("\n## Recent imports\n")
    L.append("| Finished | Supplier | Imported | Updated | Skipped | Failed "
             "| Zero-price excluded | Duplicate SKUs | Missing MPN | Rows |")
    L.append("|---|---|---:|---:|---:|---:|---:|---:|---:|---:|")
    for im in imports:
        finished = (im.get("finished_at") or "").replace("T", " ")[:16]
        L.append(f"| {finished} | {im.get('supplier_name') or im.get('supplier_code') or ''} "
                 f"| {im['imported']} | {im['updated']} | {im['skipped']} | {im['failed']} "
                 f"| {im['zero_price_excluded']} | {im['duplicate_skus']} | {im['missing_mpn']} "
                 f"| {im['total_rows']} |")
    L.append("\n_\"Cheapest\" always means lowest estimated TOTAL fulfillment cost, "
             "not lowest unit price._\n")
    path.write_text("\n".join(L), encoding="utf-8")


def _write_import_xlsx(path, coverage, imports):
    wb = Workbook()
    ws = wb.active
    ws.title = "Coverage"
    ws.append(["Metric", "Value"])
    ws.append(["Master products", coverage["total_products"]])
    ws.append(["Supplier offers", coverage["total_offers"]])
    ws.append(["Products missing mfr part #", coverage["missing_mpn"]])
    _style_header(ws); _autosize(ws)

    ws2 = wb.create_sheet("Supplier Coverage")
    ws2.append(["Supplier", "Code", "Offers", "In stock"])
    for s in coverage["per_supplier"]:
        ws2.append([s["name"], s["code"], s["n"], s["in_stock"] or 0])
    _style_header(ws2); _autosize(ws2)

    ws3 = wb.create_sheet("Import History")
    cols = ["finished_at", "supplier_name", "source_file", "imported", "updated",
            "skipped", "failed", "zero_price_excluded", "duplicate_skus",
            "missing_mpn", "total_rows"]
    ws3.append(["Finished", "Supplier", "Source file", "Imported", "Updated",
                "Skipped", "Failed", "Zero-price excluded", "Duplicate SKUs",
                "Missing MPN", "Total rows"])
    for im in imports:
        ws3.append([im.get(c) for c in cols])
    _style_header(ws3); _autosize(ws3)
    wb.save(path)


def _write_comparison_xlsx(path, comp):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric", "Value"])
    ws.append(["Jobsite", f"{comp['jobsite'].get('city')}, {comp['jobsite'].get('state')}"])
    ws.append(["Quantity per line", comp["quantity"]])
    ws.append(["Products evaluated", comp["products_evaluated"]])
    ws.append(["Products with >1 supplier", comp["multi_supplier_products"]])
    ws.append(["Products with a missing-price offer", comp["missing_price_products"]])
    ws.append(["Products with no available supplier", comp["no_available_products"]])
    _style_header(ws); _autosize(ws)

    ws2 = wb.create_sheet("Cheapest Frequency")
    ws2.append(["Supplier", "Times cheapest (by total cost)"])
    for name, n in comp["cheapest_frequency"].items():
        ws2.append([name, n])
    _style_header(ws2); _autosize(ws2)

    ws3 = wb.create_sheet("Per-Product Cheapest")
    ws3.append(["Product", "SKU", "Cheapest supplier", "Materials", "Delivery",
                "Estimated total", "Suppliers compared"])
    for r in comp["rows"][:5000]:
        ws3.append([r["product"], r["sku"], r["cheapest_supplier"],
                    r["material_subtotal"], r["delivery_cost"],
                    r["estimated_total"], r["suppliers_compared"]])
    _style_header(ws3); _autosize(ws3)
    wb.save(path)


# --------------------------------------------------------------------------
# Entry points
# --------------------------------------------------------------------------

def generate_reports(conn: sqlite3.Connection | None = None) -> list[str]:
    close = False
    if conn is None:
        conn = get_connection()
        close = True
    out = settings.REPORTS_GENERATED_DIR
    out.mkdir(parents=True, exist_ok=True)
    date = _today()

    coverage = _coverage(conn)
    imports = _latest_imports(conn)
    comp = _price_comparison(conn)

    import_md = out / f"supplier_catalog_import_{date}.md"
    import_xlsx = out / f"supplier_catalog_import_{date}.xlsx"
    comp_xlsx = out / f"supplier_price_comparison_{date}.xlsx"

    _write_import_md(import_md, coverage, imports)
    _write_import_xlsx(import_xlsx, coverage, imports)
    _write_comparison_xlsx(comp_xlsx, comp)

    if close:
        conn.close()
    return [str(import_md), str(import_xlsx), str(comp_xlsx)]


def main():
    paths = generate_reports()
    print("Generated:")
    for p in paths:
        print("  " + p)


if __name__ == "__main__":
    main()
